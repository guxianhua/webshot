import os
import sys
import json
import re
import uuid
import threading
import asyncio
import requests
import urllib3
import openpyxl
from datetime import datetime
from flask import Flask, request, jsonify, render_template, send_from_directory
from playwright.async_api import async_playwright
from PIL import Image

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# -------------------- 路径配置 --------------------
# PyInstaller 打包后，代码和模板在 sys._MEIPASS（临时解压目录）
# 数据文件放 WORK_DIR 环境变量指定的目录（用户自选工作目录）
if getattr(sys, 'frozen', False):
    BASE_DIR = sys._MEIPASS
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

WORK_DIR = os.environ.get('WORK_DIR', BASE_DIR)
TEMPLATE_DIR = os.path.join(BASE_DIR, 'templates')

UPLOAD_DIR = os.path.join(WORK_DIR, 'uploads')
SCREENSHOT_DIR = os.path.join(WORK_DIR, 'screenshots')
TASK_DIR = os.path.join(WORK_DIR, 'tasks')
DATA_FILE = os.path.join(WORK_DIR, 'sites_data.json')
META_FILE = os.path.join(WORK_DIR, 'sites_meta.json')

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(SCREENSHOT_DIR, exist_ok=True)
os.makedirs(TASK_DIR, exist_ok=True)

# Playwright 浏览器路径
# 开发模式：不设置，让 Playwright 使用系统默认路径（C:\Users\xxx\AppData\Local\ms-playwright）
# PyInstaller 打包后：强制指向内嵌的 playwright-browsers 目录
if getattr(sys, 'frozen', False):
    os.environ['PLAYWRIGHT_BROWSERS_PATH'] = os.path.join(BASE_DIR, 'playwright-browsers')

app = Flask(__name__, template_folder=TEMPLATE_DIR)

# -------------------- 启动时清理孤儿任务 --------------------
# Flask reloader 重启或进程崩溃后，可能有 task 文件残留 'running' 状态
# 启动时将这些孤儿任务标记为 'stopped'，避免前端永远显示进行中
def cleanup_orphan_tasks():
    if not os.path.exists(TASK_DIR):
        return
    for fname in os.listdir(TASK_DIR):
        if not fname.endswith('.json'):
            continue
        fpath = os.path.join(TASK_DIR, fname)
        try:
            with open(fpath, 'r', encoding='utf-8') as f:
                t = json.load(f)
            if t.get('status') == 'running':
                t['status'] = 'stopped'
                if not t.get('finished_at'):
                    t['finished_at'] = t.get('started_at') or datetime.now().isoformat()
                with open(fpath, 'w', encoding='utf-8') as f:
                    json.dump(t, f, ensure_ascii=False)
        except Exception:
            pass

cleanup_orphan_tasks()

# -------------------- 持久化 --------------------

def load_sites():
    if os.path.exists(DATA_FILE):
        try:
            with open(DATA_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            return []
    return []

def save_sites(sites):
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(sites, f, ensure_ascii=False, indent=2)

def load_meta():
    """读取网站列表元数据：{filename, edited, data_updated_at}"""
    if os.path.exists(META_FILE):
        try:
            with open(META_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            pass
    return {'filename': None, 'edited': False, 'data_updated_at': None}

def save_meta(meta):
    with open(META_FILE, 'w', encoding='utf-8') as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)

def mark_edited():
    """标记当前网站列表已被编辑，并记录更新时间"""
    meta = load_meta()
    meta['edited'] = True
    meta['data_updated_at'] = datetime.now().isoformat()
    save_meta(meta)

def touch_data_updated_at():
    """记录网站列表更新时间（上传新文件时调用）"""
    meta = load_meta()
    meta['data_updated_at'] = datetime.now().isoformat()
    save_meta(meta)

# -------------------- 工具函数 --------------------

def parse_excel(file_path):
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    websites = []
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, 2).value
        url_cell = ws.cell(row, 3).value
        unit = ws.cell(row, 4).value
        access = ws.cell(row, 5).value
        if url_cell and str(url_cell).strip():
            urls = [u.strip() for u in str(url_cell).split('\n') if u.strip()]
            for url in urls:
                if not url.startswith(('http://', 'https://')):
                    url = 'http://' + url
                safe_name = re.sub(r'[^\w\-_.]', '_', str(name) if name else 'unknown')
                safe_name = safe_name[:50]
                existing = sum(1 for w in websites if w['safe_name'].startswith(safe_name))
                if existing > 0:
                    safe_name = f"{safe_name}_{existing}"
                websites.append({
                    'id': len(websites) + 1,
                    'name': name,
                    'url': url,
                    'unit': unit,
                    'safe_name': safe_name,
                    'access': access,
                })
    return websites

def is_blank_image(img_path):
    try:
        with Image.open(img_path) as img:
            img = img.convert('RGB')
            pixels = list(img.getdata())
            total = len(pixels)
            if total == 0:
                return True
            white_pixels = sum(1 for r, g, b in pixels if r > 250 and g > 250 and b > 250)
            return white_pixels / total > 0.99
    except Exception:
        return False

# -------------------- 巡检任务 --------------------

stop_flags = {}  # task_id -> threading.Event

def run_inspect_task(task_id, websites):
    stop_flags[task_id] = threading.Event()
    task_path = os.path.join(TASK_DIR, f'{task_id}.json')
    task_data = {
        'id': task_id,
        'status': 'running',
        'total': len(websites),
        'completed': 0,
        'results': [],
        'started_at': datetime.now().isoformat(),
        'finished_at': None,
    }
    with open(task_path, 'w', encoding='utf-8') as f:
        json.dump(task_data, f, ensure_ascii=False)

    results = []
    try:
        # 阶段1+2：HTTP 探测 + 截图，统一异步并发执行
        asyncio.run(run_inspect_async(task_id, websites, results, task_data, task_path))
    except Exception as e:
        task_data['error'] = str(e)[:200]
    finally:
        # 无论成功/异常/停止，都写入终态并清理
        if stop_flags.get(task_id) and stop_flags[task_id].is_set():
            task_data['status'] = 'stopped'
        else:
            task_data['status'] = 'done'
        task_data['results'] = results
        task_data['finished_at'] = datetime.now().isoformat()
        with open(task_path, 'w', encoding='utf-8') as f:
            json.dump(task_data, f, ensure_ascii=False)
        stop_flags.pop(task_id, None)

async def run_inspect_async(task_id, websites, results, task_data, task_path):
    """统一异步执行：HTTP 探测 + 截图，使用同一事件循环和停止监听。"""
    stop_event = stop_flags.get(task_id)
    sem = asyncio.Semaphore(8)
    completed_counter = {'n': 0}  # 完成数计数器
    lock = asyncio.Lock()  # 写文件锁，避免并发写损坏

    # 预先按 idx 占位，保证 results 顺序与 websites 一致
    for site in websites:
        results.append({
            'id': site['id'],
            'name': site['name'],
            'url': site['url'],
            'unit': site.get('unit'),
            'access': site['access'],
            'http_status': None,
            'status': '无法访问',
            'error': None,
            'screenshot': None,
            'is_blank': False,
            'is_login': False,
        })

    async def flush():
        """加锁写文件，避免并发写损坏"""
        async with lock:
            with open(task_path, 'w', encoding='utf-8') as f:
                json.dump(task_data, f, ensure_ascii=False)

    async def probe_and_shoot(idx, site):
        async with sem:
            if stop_event and stop_event.is_set():
                completed_counter['n'] += 1
                task_data['completed'] = completed_counter['n']
                await flush()
                return
            url = site['url']
            result = results[idx]

            # HTTP 探测（异步，避免阻塞事件循环）
            if not (stop_event and stop_event.is_set()):
                try:
                    resp = await asyncio.to_thread(
                        requests.get, url, timeout=8, verify=False, allow_redirects=True
                    )
                    result['http_status'] = resp.status_code
                    result['status'] = '正常' if resp.status_code < 400 else '异常'
                except Exception as e:
                    result['error'] = str(e)[:100]
                    result['status'] = '无法访问'

            # 截图
            if not (stop_event and stop_event.is_set()):
                await take_one_screenshot(browser, task_id, site, result)

            completed_counter['n'] += 1
            task_data['completed'] = completed_counter['n']
            await flush()

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        tasks = [asyncio.create_task(probe_and_shoot(i, site)) for i, site in enumerate(websites)]

        async def watch_stop():
            while True:
                if stop_event and stop_event.is_set():
                    for t in tasks:
                        if not t.done():
                            t.cancel()
                    return
                await asyncio.sleep(0.5)

        watcher = asyncio.create_task(watch_stop())
        try:
            await asyncio.gather(*tasks, return_exceptions=True)
        finally:
            if not watcher.done():
                watcher.cancel()
                try:
                    await watcher
                except asyncio.CancelledError:
                    pass
            await browser.close()

async def take_one_screenshot(browser, task_id, site, result):
    """对单个网站截图并更新 result。失败则重试一次，仍失败则标记为"正常但无截图"。"""
    filename = f"{task_id}_{site['id']:03d}_{site['safe_name']}.png"
    filepath = os.path.join(SCREENSHOT_DIR, filename)

    async def _try(wait_until, goto_timeout, sleep_sec, shot_timeout):
        page = await browser.new_page(viewport={'width': 1920, 'height': 1080})
        try:
            await asyncio.wait_for(
                page.goto(site['url'], wait_until=wait_until, timeout=goto_timeout),
                timeout=float(goto_timeout) / 1000 + 5
            )
            if sleep_sec:
                await asyncio.sleep(sleep_sec)
            await asyncio.wait_for(
                page.screenshot(path=filepath, full_page=True),
                timeout=float(shot_timeout) / 1000 + 2
            )
            result['screenshot'] = filename
            result['is_blank'] = is_blank_image(filepath)
            current_url = page.url
            page_text = await page.content()
            is_login = 'auth.shsmu.edu.cn' in current_url or '忘记密码' in page_text
            result['is_login'] = is_login
            if is_login:
                result['status'] = '登录页面'
            elif result['status'] == '无法访问' or result['status'] == '异常':
                result['status'] = '正常'
            return True
        except asyncio.CancelledError:
            raise
        except Exception as e:
            result['error'] = str(e)[:100]
            return False
        finally:
            try:
                await page.close()
            except Exception:
                pass

    # 第一次：标准参数
    ok = await _try('domcontentloaded', 15000, 1.5, 10000)
    # 第二次重试：用更宽松的条件（不等 DOM，快速截图）
    if not ok and (result.get('status') == '正常' or result.get('status') == '登录页面'):
        ok = await _try('commit', 8000, 0.5, 8000)
    # 仍然失败且状态是正常 → 标记为正常但无截图
    if not ok and result.get('status') == '正常':
        result['status'] = '正常但无截图'
        result['is_no_screenshot'] = True

# -------------------- 路由 --------------------

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/admin')
def admin():
    return render_template('admin.html')

@app.route('/user')
def user():
    return render_template('user.html')

@app.route('/api/upload', methods=['POST'])
def api_upload():
    file = request.files.get('file')
    if not file:
        return jsonify({'error': '未上传文件'}), 400
    ext = os.path.splitext(file.filename)[1].lower()
    if ext != '.xlsx':
        return jsonify({'error': '只允许上传 .xlsx 文件'}), 400
    task_id = str(uuid.uuid4())
    save_path = os.path.join(UPLOAD_DIR, f'{task_id}{ext}')
    file.save(save_path)
    try:
        websites = parse_excel(save_path)
        save_sites(websites)
        # 保存原始文件名，记录更新时间
        now = datetime.now().isoformat()
        save_meta({'filename': file.filename, 'edited': False, 'data_updated_at': now})
        return jsonify({'task_id': task_id, 'websites': websites})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/sites')
def api_sites():
    sites = load_sites()
    return jsonify({'sites': sites})

@app.route('/api/sites/meta')
def api_sites_meta():
    return jsonify(load_meta())

@app.route('/api/sites/save', methods=['POST'])
def api_save_sites():
    data = request.get_json()
    sites = data.get('sites', [])
    save_sites(sites)
    mark_edited()
    return jsonify({'success': True})

@app.route('/api/sites/delete/<int:site_id>', methods=['DELETE'])
def api_delete_site(site_id):
    sites = load_sites()
    sites = [s for s in sites if s.get('id') != site_id]
    save_sites(sites)
    mark_edited()
    return jsonify({'success': True})

@app.route('/api/inspect', methods=['POST'])
def api_inspect():
    data = request.get_json()
    task_id = data.get('task_id')
    websites = data.get('websites')
    if not websites:
        websites = load_sites()
        if not websites:
            return jsonify({'error': '请先在管理端上传网站列表'}), 400
    inspect_id = str(uuid.uuid4())
    t = threading.Thread(target=run_inspect_task, args=(inspect_id, websites), daemon=True)
    t.start()
    return jsonify({'inspect_id': inspect_id})

@app.route('/api/inspect/stop', methods=['POST'])
def api_stop_inspect():
    data = request.get_json() or {}
    task_id = data.get('task_id')
    # 1. 停止正在运行的线程任务
    if task_id and task_id in stop_flags:
        stop_flags[task_id].set()
        return jsonify({'success': True})
    for tid, ev in stop_flags.items():
        ev.set()
    # 2. 处理孤儿任务：task 文件显示 running 但当前进程无对应线程
    #    （Flask reloader 重启或进程崩溃后残留的任务）
    if os.path.exists(TASK_DIR):
        for fname in os.listdir(TASK_DIR):
            if not fname.endswith('.json'):
                continue
            fpath = os.path.join(TASK_DIR, fname)
            try:
                with open(fpath, 'r', encoding='utf-8') as f:
                    t = json.load(f)
                if t.get('status') == 'running':
                    t['status'] = 'stopped'
                    if not t.get('finished_at'):
                        t['finished_at'] = datetime.now().isoformat()
                    with open(fpath, 'w', encoding='utf-8') as f:
                        json.dump(t, f, ensure_ascii=False)
            except Exception:
                pass
    return jsonify({'success': True})

@app.route('/api/task/<inspect_id>')
def api_task(inspect_id):
    task_path = os.path.join(TASK_DIR, f'{inspect_id}.json')
    if not os.path.exists(task_path):
        return jsonify({'error': '任务不存在'}), 404
    with open(task_path, 'r', encoding='utf-8') as f:
        return jsonify(json.load(f))

@app.route('/api/screenshot/<filename>')
def api_screenshot(filename):
    return send_from_directory(SCREENSHOT_DIR, filename)

@app.route('/api/latest-task')
def api_latest_task():
    """返回最近一次巡检任务的状态（task）以及用于相册展示的任务（album_task）。
    - 最新任务正在运行：album_task = 最新任务（显示当前巡检的部分结果）
    - 最新任务已结束且有结果：album_task = 最新任务
    - 最新任务已结束但无结果（如刚启动即停止）：album_task 回退为上一个有结果的任务
    - 如果网站列表在任务完成后被更新过（上传/编辑/删除），album_task 返回 None"""
    if not os.path.exists(TASK_DIR):
        return jsonify({'task': None, 'album_task': None})
    task_files = [f for f in os.listdir(TASK_DIR) if f.endswith('.json')]
    if not task_files:
        return jsonify({'task': None, 'album_task': None})
    task_files.sort(key=lambda f: os.path.getmtime(os.path.join(TASK_DIR, f)), reverse=True)
    task = None
    album_task = None
    for f in task_files:
        with open(os.path.join(TASK_DIR, f), 'r', encoding='utf-8') as fp:
            t = json.load(fp)
        if task is None:
            task = t
            # 最新任务正在运行或已有结果，直接用作相册数据
            if t.get('status') == 'running' or t.get('results'):
                album_task = t
                break
            # 最新任务已结束但无结果，继续找上一个有结果的
            continue
        if t.get('results'):
            album_task = t
            break
    # 如果任务早于网站列表更新时间（非运行中），标记过期
    if task and task.get('status') != 'running':
        meta = load_meta()
        data_updated_at = meta.get('data_updated_at')
        if data_updated_at and task.get('started_at'):
            try:
                started = datetime.fromisoformat(task['started_at'])
                updated = datetime.fromisoformat(data_updated_at)
                if updated > started:
                    if isinstance(task, dict):
                        task['outdated'] = True
            except Exception:
                pass
    if album_task and album_task.get('status') != 'running':
        meta = load_meta()
        data_updated_at = meta.get('data_updated_at')
        if data_updated_at and album_task.get('started_at'):
            try:
                started = datetime.fromisoformat(album_task['started_at'])
                updated = datetime.fromisoformat(data_updated_at)
                if updated > started:
                    album_task = None
            except Exception:
                pass
    return jsonify({'task': task, 'album_task': album_task})

if __name__ == '__main__':
    port = int(os.environ.get('FLASK_PORT', '5000'))
    use_debug = os.environ.get('FLASK_DEBUG', '1') == '1'
    app.run(host='0.0.0.0', port=port, debug=use_debug)
